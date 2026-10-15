# Aspa — multiple-choice exams: generate, grade and review
# Copyright (C) 2026 Lorena Gutiérrez Madroñal, Rubén Baena Pérez, José Miguel Mota Macías
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Lectura del banco de preguntas y de la lista de alumnos.

Acepta:
  * El formato nuevo (plantilla_preguntas.xlsx): id | tema | pregunta | A | B | C | D | correcta | grupo | fija
  * Cabeceras alternativas habituales: num | Tema | Preg | Resp 1..4 | Corr | Selec
  * CSV con las mismas cabeceras.
La respuesta correcta puede escribirse con letras ("AC") o con números como antes ("13").
"""
from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .i18n import t

LETRAS = "ABCD"

# Opciones que deben quedarse en su sitio al desordenar ("Todas son correctas", etc.)
PATRON_FIJA = re.compile(
    r"^\s*(todas|ninguna|ambas|las dos|ninguna de las|todas las anteriores"
    r"|all of|all are|all the|none of|none is|none are|both|neither)\b", re.I
)


def _norm(s) -> str:
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@dataclass
class Pregunta:
    id: str
    texto: str
    opciones: list[str]
    correctas: set[int]           # índices 0-3 en el orden original
    grupo: str = ""
    tema: str = ""
    fijas: set[int] = field(default_factory=set)
    fila: int = 0
    explicacion: str = ""         # opcional: se muestra en la revisión si el alumno no acierta

    def to_dict(self):
        return {
            "id": self.id, "texto": self.texto, "opciones": self.opciones,
            "correctas": "".join(LETRAS[i] for i in sorted(self.correctas)),
            "grupo": self.grupo, "tema": self.tema, "explicacion": self.explicacion,
        }


def parse_correctas(v) -> set[int]:
    s = _texto(v).upper().replace(" ", "").replace(",", "").replace(";", "")
    out = set()
    for ch in s:
        if ch in LETRAS:
            out.add(LETRAS.index(ch))
        elif ch in "1234":
            out.add(int(ch) - 1)
    return out


# Sinónimos de cabecera -> campo interno
CABECERAS = {
    # sinónimos en español y en inglés (las plantillas se generan en el idioma de Aspa)
    "id": ["id", "num", "numero", "n", "codigo", "number", "code"],
    "tema": ["tema", "topic", "unit"],
    "texto": ["pregunta", "preg", "enunciado", "question", "stem"],
    "A": ["a", "resp1", "respuesta1", "opciona", "optiona", "answer1"],
    "B": ["b", "resp2", "respuesta2", "opcionb", "optionb", "answer2"],
    "C": ["c", "resp3", "respuesta3", "opcionc", "optionc", "answer3"],
    "D": ["d", "resp4", "respuesta4", "opciond", "optiond", "answer4"],
    "correcta": ["correcta", "correctas", "corr", "solucion", "correct", "answer", "key"],
    "grupo": ["grupo", "selec", "select", "seleccion", "sel", "group"],
    "fija": ["fija", "fijas", "nodesordenar", "fixed", "keep"],
    "explicacion": ["explicacion", "explicaciones", "justificacion", "retroalimentacion", "comentario",
                    "explanation", "feedback", "rationale"],
}


def _mapear_cabecera(fila) -> dict[str, int]:
    pos = {}
    normas = [_norm(c) for c in fila]
    for campo, sinonimos in CABECERAS.items():
        for i, n in enumerate(normas):
            if n in sinonimos and campo not in pos:
                pos[campo] = i
                break
    return pos


def _filas_tabla(ruta: Path, hojas_preferidas: list[str]):
    """Devuelve las filas (listas) de la primera hoja que coincida."""
    ruta = Path(ruta)
    if ruta.suffix.lower() in (".csv", ".tsv", ".txt"):
        texto = ruta.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(texto[:4000], delimiters=",;\t")
        return list(csv.reader(texto.splitlines(), dialect))
    wb = load_workbook(ruta, read_only=True, data_only=True)
    nombres = {_norm(n): n for n in wb.sheetnames}
    hoja = None
    for pref in hojas_preferidas:
        if _norm(pref) in nombres:
            hoja = wb[nombres[_norm(pref)]]
            break
    if hoja is None:
        if hojas_preferidas and hojas_preferidas[0] == "__estricto__":
            return None
        hoja = wb[wb.sheetnames[0]]
    filas = [list(r) for r in hoja.iter_rows(values_only=True)]
    wb.close()
    return filas


def cargar_preguntas(ruta: Path) -> tuple[list[Pregunta], list[str]]:
    """Devuelve (preguntas, avisos)."""
    filas = _filas_tabla(ruta, ["preguntas", "banco", "questions", "bank"])
    avisos: list[str] = []
    cab_idx, pos = None, {}
    for i, f in enumerate(filas[:15]):
        p = _mapear_cabecera(f)
        if "texto" in p and "A" in p:
            cab_idx, pos = i, p
            break
    if cab_idx is None:
        raise ValueError(t("No encuentro la fila de cabeceras (se esperan columnas como "
                           "'pregunta', 'A', 'B', 'C', 'D', 'correcta')."))
    preguntas, ids = [], set()
    for n, f in enumerate(filas[cab_idx + 1:], start=cab_idx + 2):
        def col(c):
            return _texto(f[pos[c]]) if c in pos and pos[c] < len(f) else ""
        texto = col("texto")
        if not texto:
            continue
        opciones = [col(L) for L in LETRAS]
        qid = col("id") or str(n - cab_idx - 1)
        if qid in ids:
            avisos.append(t("Fila {n}: el identificador {id} está repetido; se renombra a {nuevo}.", n=n, id=qid, nuevo=f"{qid}_{n}"))
            qid = f"{qid}_{n}"
        ids.add(qid)
        correctas = parse_correctas(col("correcta"))
        if not correctas:
            avisos.append(t("Pregunta {id} (fila {n}): sin respuesta correcta.", id=qid, n=n))
        for c in list(correctas):
            if not opciones[c]:
                avisos.append(t("Pregunta {id}: la respuesta correcta {letra} está vacía.", id=qid, letra=LETRAS[c]))
        fijas = {i for i, o in enumerate(opciones) if o and PATRON_FIJA.match(o)}
        fijas |= parse_correctas(col("fija"))
        preguntas.append(Pregunta(
            id=qid, texto=texto, opciones=opciones, correctas=correctas,
            grupo=col("grupo"), tema=col("tema"), fijas=fijas, fila=n,
            explicacion=col("explicacion"),
        ))
    if not preguntas:
        raise ValueError(t("El fichero no contiene preguntas."))
    return preguntas, avisos


# ---------------------------------------------------------------- alumnos

@dataclass
class Alumno:
    documento: str
    nombre: str
    superadas: list[str] = field(default_factory=list)  # ids de preguntas ya superadas en clase

    @property
    def digitos(self) -> str:
        return re.sub(r"\D", "", self.documento)


_DOC = ("doc", "dni", "documento", "nif", "nie", "pasaporte", "identificador", "id",
        "document", "iddocument", "passport", "studentid")
_NOMBRE_COMPLETO = ("apellidosynombre", "apellidosnombre", "nombrecompleto", "alumno", "estudiante",
                    "fullname", "student", "name")
_PARTES = ("apellido1", "primerapellido", "apellidos", "apellido", "apellido2", "segundoapellido", "nombre",
           "surname", "lastname", "firstname")
_SUPERADAS = ("superadas", "aprobadas", "eliminadas", "liberadas", "preguntassuperadas", "exentas",
              "passed", "exempt", "removed", "passedquestions")


def parse_superadas(v) -> list[str]:
    """«2, 5; 7 12» -> ['2', '5', '7', '12']"""
    out = []
    for trozo in re.split(r"[,;\s]+", _texto(v)):
        trozo = trozo.strip()
        if trozo and trozo not in out:
            out.append(trozo)
    return out


def _entero(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def cargar_alumnos(ruta: Path) -> list[Alumno]:
    """Lee la lista de clase.

    Admite la lista de clase que exporta la universidad, una hoja simple «documento | nombre | superadas» y la hoja «nombres»
    de exámenes personalizados con columnas sin cabecera (nombre en «Apellido 1» y los números de las
    preguntas superadas en las columnas siguientes sin cabecera, hasta 15).
    """
    filas = _filas_tabla(ruta, ["alumnos", "nombres", "lista", "estudiantes", "students", "names", "list"])
    if filas is None:
        return []
    cab = doc_c = None
    for i, f in enumerate(filas[:15]):
        normas = [_norm(c) for c in f]
        d = next((j for j, n in enumerate(normas) if n in _DOC), None)
        tiene_nombre = any(n in _NOMBRE_COMPLETO or n in _PARTES for n in normas)
        if d is not None and tiene_nombre:
            cab, doc_c = i, d
            break
    if cab is None:
        raise ValueError(t("No encuentro las columnas de documento y nombre en la lista de alumnos."))
    normas = [_norm(c) for c in filas[cab]]
    completo = next((j for j, n in enumerate(normas) if n in _NOMBRE_COMPLETO), None)
    partes = [j for j, n in enumerate(normas) if n in _PARTES]
    nombre_cols = [completo] if completo is not None else partes
    sup_c = next((j for j, n in enumerate(normas) if n in _SUPERADAS), None)

    # Formato antiguo: columnas sin cabecera justo detrás del nombre, con números de pregunta
    antiguas = []
    if sup_c is None and nombre_cols:
        j = max(nombre_cols) + 1
        while j < len(normas) and not normas[j] and len(antiguas) < 15:
            antiguas.append(j)
            j += 1
        con_doc = [f for f in filas[cab + 1:] if doc_c < len(f) and _texto(f[doc_c])]
        valores = [f[k] for f in con_doc for k in antiguas if k < len(f) and f[k] not in (None, "")]
        if not valores or sum(_entero(v) is None for v in valores) > 0.1 * len(valores):
            antiguas = []

    alumnos = []
    for f in filas[cab + 1:]:
        if doc_c >= len(f):
            continue
        doc = _texto(f[doc_c])
        if not doc:
            continue
        trozos = [_texto(f[k]) for k in nombre_cols if k < len(f) and _texto(f[k])]
        if completo is None and len(trozos) > 1 and "," not in " ".join(trozos):
            # «Apellido1 Apellido2, Nombre» si vienen separados
            nom = " ".join(trozos[:-1]) + ", " + trozos[-1]
        else:
            nom = " ".join(trozos)
        if sup_c is not None:
            sup = parse_superadas(f[sup_c]) if sup_c < len(f) else []
        else:
            sup = []
            for k in antiguas:
                n = _entero(f[k]) if k < len(f) else None
                if n is not None and str(n) not in sup:
                    sup.append(str(n))
        alumnos.append(Alumno(doc, nom, sup))
    return alumnos


def buscar_alumno(leido: str, alumnos: list[Alumno]):
    """Empareja el identificador leído (dígitos, '?' = dudoso) con la lista.

    Devuelve (alumno | None, tipo) con tipo en:
      'exacto', 'comodin' (había dígitos dudosos), 'aproximado' (1 dígito distinto),
      'no_encontrado', 'vacio'.
    """
    s = (leido or "").strip()
    if not s.replace("?", ""):
        return None, "vacio"

    def variantes(d):
        return {d, d.lstrip("0")}

    def coincide(patron, d):
        return len(patron) == len(d) and all(p == "?" or p == c for p, c in zip(patron, d))

    for pat in variantes(s):
        cand = [a for a in alumnos if any(coincide(pat, v) for v in variantes(a.digitos))]
        if len(cand) == 1:
            return cand[0], ("comodin" if "?" in pat else "exacto")
    # un dígito distinto
    cand = []
    for a in alumnos:
        for v in variantes(a.digitos):
            for pat in variantes(s):
                if len(v) == len(pat) and sum(p != c and p != "?" for p, c in zip(pat, v)) == 1:
                    cand.append(a)
    uniq = {a.documento: a for a in cand}
    if len(uniq) == 1:
        return next(iter(uniq.values())), "aproximado"
    return None, "no_encontrado"
