# Aspa — multiple-choice exams: generate, grade and review
# Copyright (C) 2026 Lorena Gutiérrez Madroñal, Rubén Baena Pérez, José Miguel Mota Macías
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Plantillas Excel para el banco de preguntas y la lista de alumnos.

Cada plantilla existe en dos versiones:
  * limpia: solo las cabeceras, lista para rellenar.
  * ejemplo: con filas de muestra que ilustran cada opción y una hoja de instrucciones.

Los ejemplos se definen aquí una sola vez y se reutilizan en las páginas de ayuda, de modo
que lo que se ve en pantalla y lo que se descarga coincide siempre.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .i18n import t

AZUL = "1F3A93"
FUENTE = "Arial"

COLUMNAS_PREGUNTAS = [
    ("id", 7, "Identificador único de la pregunta (número o texto). Se usa en la estadística."),
    ("tema", 14, "Opcional. Solo informativo."),
    ("pregunta", 58, "Enunciado de la pregunta."),
    ("A", 30, "Opción A."),
    ("B", 30, "Opción B."),
    ("C", 30, "Opción C (puede quedar vacía si la pregunta tiene menos opciones)."),
    ("D", 30, "Opción D (puede quedar vacía)."),
    ("correcta", 10, "Letra o letras correctas: B, AC, ABD… (también vale 2, 13…)."),
    ("grupo", 8, "Código para asignar la pregunta a modelos. Vacío = pregunta de relleno al azar."),
    ("fija", 7, "Letras que no se mueven al desordenar (opcional)."),
    ("explicacion", 45, "Opcional. Explicación de la respuesta correcta: se muestra en la revisión del "
                        "examen cuando el alumno no ha acertado la pregunta."),
]

# Cada ejemplo: (fila, explicación para la ayuda)
EJEMPLOS_PREGUNTAS = [
    ({"id": "1", "tema": "Color", "pregunta": "¿Cuántos canales tiene el modelo de color RGB?",
      "A": "Uno", "B": "Dos", "C": "Tres", "D": "Cuatro", "correcta": "C", "grupo": "1",
      "explicacion": "RGB tiene tres canales: rojo, verde y azul."},
     "Una sola respuesta correcta. Pertenece al grupo 1. Con explicación: se verá en la revisión si el alumno la falla."),
    ({"id": "2", "tema": "Color", "pregunta": "¿Qué modelos de color son sustractivos? (una o más respuestas)",
      "A": "CMYK", "B": "RYB", "C": "RGB", "D": "HSL", "correcta": "AB", "grupo": "1",
      "explicacion": "CMYK y RYB mezclan pigmentos (sustractivos); RGB mezcla luz (aditivo) y HSL es otra forma de describir el color."},
     "Varias respuestas correctas: se escriben juntas (AB). Conviene indicarlo en el enunciado."),
    ({"id": "3", "tema": "Imagen", "pregunta": "Un píxel… (una o más respuestas)",
      "A": "Es uno de los componentes de una imagen digital",
      "B": "Tiene un valor según la combinación de colores del modelo",
      "C": "Equivale siempre a un bit", "D": "Todas son correctas", "correcta": "AB", "grupo": "2"},
     "La opción «Todas son correctas» se queda siempre en la D al desordenar (es automático)."),
    ({"id": "4", "tema": "Imagen", "pregunta": "¿Qué formato de imagen admite transparencia?",
      "A": "PNG", "B": "JPEG", "C": "BMP", "D": "Ninguna es correcta", "correcta": "1", "grupo": "2",
      "explicacion": "PNG admite canal alfa (transparencia); JPEG y BMP no."},
     "La respuesta correcta también puede escribirse con números (1 = A, 2 = B…). «Ninguna…» se queda fija en la D."),
    ({"id": "5", "tema": "Formatos", "pregunta": "El formato SVG es vectorial.",
      "A": "Verdadero", "B": "Falso", "C": "", "D": "", "correcta": "A", "grupo": ""},
     "Pregunta de verdadero/falso: se dejan vacías C y D. Sin grupo: sirve de relleno al azar."),
    ({"id": "6", "tema": "Resolución", "pregunta": "Ordena de menor a mayor resolución habitual:",
      "A": "72 ppp", "B": "150 ppp", "C": "300 ppp", "D": "Ninguna es correcta",
      "correcta": "D", "grupo": "", "fija": "ABC"},
     "Columna «fija» = ABC: esas opciones no se desordenan (útil cuando el orden importa)."),
    ({"id": "7", "tema": "Compresión", "pregunta": "La compresión JPEG…",
      "A": "Es con pérdida", "B": "Es sin pérdida", "C": "No existe", "D": "Solo sirve para vídeo",
      "correcta": "A", "grupo": ""},
     "Sin grupo: puede salir al azar en cualquier modelo que necesite completar preguntas."),
]

EJEMPLOS_ALUMNOS = [
    ({"documento": "12345678Z", "nombre": "García López, Ana", "superadas": "1, 2"},
     "DNI: se comparan solo los números (12345678). En un personalizado no tendrá las preguntas 1 y 2."),
    ({"documento": "X1234567L", "nombre": "Dubois, Camille", "superadas": ""},
     "NIE: también solo los números (1234567). Sin superadas: examen completo."),
    ({"documento": "YA7731982", "nombre": "Rossi, Giulia", "superadas": "3"},
     "Pasaporte con letras: se usan sus números (7731982). En un personalizado no tendrá la 3."),
    ({"documento": "C01X00T47", "nombre": "Müller, Lena", "superadas": ""},
     "Pasaporte con letras intercaladas: números 010047. El alumno marca solo esos dígitos."),
]

INSTRUCCIONES_PREGUNTAS = [
    "Cómo rellenar el banco de preguntas",
    "",
    "Una fila por pregunta en la hoja «preguntas». Las columnas son:",
    "__COLUMNAS__",
    "",
    "Varias respuestas correctas: escriba todas las letras juntas en «correcta» (por ejemplo AC).",
    "Con puntuación parcial, quien marque solo una parte de las correctas recibe parte del punto.",
    "",
    "Grupos y modelos:",
    "  En la aplicación, cada modelo (A, B, C…) indica qué grupos usa, por ejemplo «1,2».",
    "  El modelo toma todas las preguntas de esos grupos. Si no llegan al número de preguntas",
    "  del examen, el resto se elige al azar entre las preguntas SIN grupo (distintas en cada modelo).",
    "  Un modelo sin grupos toma todas sus preguntas al azar de todo el banco.",
    "",
    "Desorden:",
    "  Las preguntas y las opciones se desordenan en cada modelo (se puede desactivar).",
    "  Las opciones que empiezan por «Todas…» o «Ninguna…» se quedan siempre en su sitio.",
    "  Para fijar otras, escriba sus letras en la columna «fija».",
    "",
    "Menos de cuatro opciones: deje vacías las últimas columnas (por ejemplo, C y D en verdadero/falso).",
    "",
    "Explicaciones (opcional): en la columna «explicacion» puede escribir por qué la respuesta correcta lo es.",
    "Se muestra en la pantalla de revisión, bajo la pregunta, cuando el alumno no la ha acertado (fallada,",
    "parcial o en blanco). Es útil en la revisión con el alumno. Se pueden añadir o cambiar después de",
    "generar el examen: basta con volver a subir el banco.",
]

INSTRUCCIONES_ALUMNOS = [
    "Cómo rellenar la lista de alumnos",
    "",
    "Dos columnas en la hoja «alumnos»: documento y nombre.",
    "  documento: DNI, NIE o pasaporte, tal cual. Se puede escribir con letras.",
    "  nombre: preferiblemente «Apellidos, Nombre», que es como aparecerá en las notas.",
    "",
    "Para identificar al alumno se comparan solo los NÚMEROS del documento con lo que marca en la",
    "hoja de respuestas (10 casillas). Por eso funciona con NIE y pasaportes.",
    "",
    "También puede subir directamente, sin modificarla, la lista de clase que exporta el sistema de gestión académica:",
    "la aplicación localiza las columnas «Doc.» y «Apellidos y nombre».",
    "",
    "Columna «superadas» (solo para exámenes personalizados): números (id) de las preguntas del banco",
    "que el alumno ya ha superado en clase, separados por comas. No aparecerán en su examen y cuentan",
    "como acierto. Déjela vacía si no tiene ninguna. En los exámenes por modelos se ignora.",
]


INSTRUCCIONES_PERSONALIZADO = [
    "Examen personalizado: preguntas y alumnos en un solo fichero",
    "",
    "Este Excel tiene dos hojas. Súbalo entero en «Banco de preguntas» (paso 1): Aspa lee las",
    "preguntas de la hoja «preguntas» y la lista de alumnos de la hoja «alumnos». No hace falta",
    "subir la lista aparte (si se sube, tiene prioridad).",
    "",
    "Hoja «preguntas»: igual que el banco normal (id, tema, pregunta, A-D, correcta, grupo, fija, explicacion).",
    "  El «id» es el número con el que se indican las preguntas superadas.",
    "  El «grupo» sirve para formar los modelos; cada alumno recibe un modelo que contenga",
    "  todas sus superadas.",
    "",
    "Hoja «alumnos»: documento, nombre y superadas.",
    "  superadas: ids de las preguntas que el alumno ya ha superado en clase, separados por comas",
    "  (por ejemplo 1, 2). Se quitan de su examen y cada una cuenta como un acierto completo.",
    "  Vacía = examen completo con su nombre impreso.",
    "",
    "La nota se calcula sobre el examen completo (las preguntas por examen del paso 2) y se pasa a 10.",
    "Para quien no esté en la lista se pueden imprimir exámenes de reserva completos.",
]


def _estilo_cabecera(fila):
    for c in fila:
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center")


# ------------------------------------------------------------------ versiones traducidas
# Los textos de este módulo están en español y se traducen al idioma actual al generar
# cada plantilla (las cabeceras traducidas también se reconocen al leer el fichero).

_TRADUCIBLES = ("tema", "pregunta", "A", "B", "C", "D", "explicacion")


def columnas() -> list[tuple[str, int, str]]:
    return [(t(c), ancho, t(desc)) for c, ancho, desc in COLUMNAS_PREGUNTAS]


def ejemplos_preguntas() -> list:
    return [({k: (t(v) if k in _TRADUCIBLES and v else v) for k, v in fila.items()}, t(expl))
            for fila, expl in EJEMPLOS_PREGUNTAS]


def ejemplos_alumnos() -> list:
    return [(dict(fila), t(expl)) for fila, expl in EJEMPLOS_ALUMNOS]


def _hoja_instrucciones(wb, lineas):
    ws = wb.create_sheet(t("instrucciones"), 0)
    for linea in lineas:
        if linea == "__COLUMNAS__":
            for c, _, d in columnas():
                ws.append([f"  {c}: {d}"])
        else:
            ws.append([t(linea) if linea.strip() else ""])
    ws["A1"].font = Font(name=FUENTE, bold=True, size=13)
    for fila in ws.iter_rows(min_row=2):
        fila[0].font = Font(name=FUENTE)
    ws.column_dimensions["A"].width = 105
    return ws


def _a_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hoja_preguntas(ws, con_ejemplos: bool):
    ws.title = t("preguntas")
    claves = [c for c, _, _ in COLUMNAS_PREGUNTAS]
    ws.append([c for c, _, _ in columnas()])
    _estilo_cabecera(ws[1])
    for i, (c, ancho, desc) in enumerate(columnas(), 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = ancho
        ws.cell(1, i).comment = Comment(desc, "Aspa")
    if con_ejemplos:
        for fila, _ in ejemplos_preguntas():
            ws.append([fila.get(c, "") for c in claves])
    borde = Border(bottom=Side(style="thin", color="D6DCE6"))
    for fila in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2)):
        for c in fila:
            c.font = Font(name=FUENTE)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = borde
    # columnas de texto para que Excel no convierta «13» o «1,2» en números
    for col in ("A", "H", "I", "J"):
        for r in range(2, 501):
            ws[f"{col}{r}"].number_format = "@"
    dv = DataValidation(type="custom", formula1='=OR(H2="",ISNUMBER(SEARCH(LEFT(H2,1),"ABCD1234")))',
                        showErrorMessage=True, errorTitle=t("Respuesta correcta"),
                        error=t("Escriba letras A-D (por ejemplo AC) o números 1-4."))
    ws.add_data_validation(dv)
    dv.add("H2:H500")
    ws.freeze_panes = "A2"


def _hoja_alumnos(ws, con_ejemplos: bool):
    ws.title = t("alumnos")
    ws.append([t("documento"), t("nombre"), t("superadas")])
    _estilo_cabecera(ws[1])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 34
    ws["C1"].comment = Comment(t("Solo exámenes personalizados: números (id) de las preguntas que el alumno "
                                 "ya ha superado y se quitan de su examen, separados por comas (p. ej. 2, 5, 7)."),
                               "Aspa")
    if con_ejemplos:
        for fila, _ in EJEMPLOS_ALUMNOS:
            ws.append([fila["documento"], fila["nombre"], fila["superadas"]])
    for r in range(2, 501):
        ws[f"A{r}"].number_format = "@"  # que Excel no quite ceros a la izquierda
        ws[f"C{r}"].number_format = "@"  # «2, 5» no debe convertirse en número
        for c in ("A", "B", "C"):
            ws[f"{c}{r}"].font = Font(name=FUENTE)
    ws.freeze_panes = "A2"


def plantilla_preguntas(con_ejemplos: bool) -> bytes:
    wb = Workbook()
    _hoja_preguntas(wb.active, con_ejemplos)
    _hoja_instrucciones(wb, INSTRUCCIONES_PREGUNTAS)
    wb.active = 1
    return _a_bytes(wb)


def plantilla_alumnos(con_ejemplos: bool) -> bytes:
    wb = Workbook()
    _hoja_alumnos(wb.active, con_ejemplos)
    _hoja_instrucciones(wb, INSTRUCCIONES_ALUMNOS)
    wb.active = 1
    return _a_bytes(wb)


def plantilla_personalizado(con_ejemplos: bool) -> bytes:
    """Un único Excel con las hojas «preguntas» y «alumnos» (con la columna «superadas»).
    Se sube entero en «Banco de preguntas» y Aspa toma la lista de alumnos de la misma hoja."""
    wb = Workbook()
    _hoja_preguntas(wb.active, con_ejemplos)
    _hoja_alumnos(wb.create_sheet("alumnos"), con_ejemplos)
    _hoja_instrucciones(wb, INSTRUCCIONES_PERSONALIZADO)
    wb.active = 1
    return _a_bytes(wb)
