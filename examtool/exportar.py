# Aspa — multiple-choice exams: generate, grade and review
# Copyright (C) 2026 Lorena Gutiérrez Madroñal, Rubén Baena Pérez, José Miguel Mota Macías
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Exportación de resultados a Excel.

Es una «foto» de las notas en el momento de exportar: los valores se escriben ya
calculados (la fuente de verdad es la aplicación, donde se recalcula todo).
"""
from __future__ import annotations

import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calificar import calificar_todo, estadistica
from .i18n import t

AZUL = "1F3A93"
FUENTE = "Arial"


def _cabecera(ws, titulos, anchos=None):
    ws.append([t(x) for x in titulos])  # los títulos se escriben en español y se traducen aquí
    for c in ws[ws.max_row]:
        c.font = Font(name=FUENTE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = f"A{ws.max_row + 1}"
    for i, w in enumerate(anchos or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fuente(ws):
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if c.font.name != FUENTE:
                c.font = Font(name=FUENTE, bold=c.font.bold, color=c.font.color)


def exportar_excel(proyecto, destino=None) -> str:
    cfg = proyecto.config()
    gen = proyecto.generado()
    res = [r for r in calificar_todo(proyecto) if r.get("estado") != "vacia"]
    alumnos = proyecto.alumnos()
    wb = Workbook()

    # ---------------------------------------------------------------- Notas
    ws = wb.active
    ws.title = t("Notas")
    _cabecera(ws, ["Documento", "Apellidos y nombre", t("Nota (sobre {n})", n=f"{cfg['nota_maxima']:g}"), "Puntos",
                   "Bien", "Parcial", "Mal", "En blanco", "Por revisar", "Superadas", "Modelo", "Copia",
                   "Archivo", "Incidencias"],
              [14, 34, 11, 8, 6, 7, 6, 8, 9, 9, 8, 7, 30, 60])
    amarillo = PatternFill("solid", fgColor="FFF2B3")
    for r in sorted(res, key=lambda r: ((r.get("alumno") or {}).get("nombre") or "~", r["clave"] or "")):
        a = r.get("alumno") or {}
        c = r.get("cuenta", {})
        ws.append([a.get("documento", r.get("id_leido", "")), a.get("nombre", ""), r.get("nota"),
                   None if r.get("puntos") is None else round(r["puntos"], 2), c.get("bien"), c.get("parcial"), c.get("mal"), c.get("blanco"),
                   c.get("revisar"), c.get("superadas", 0), r.get("modelo_nombre", ""), r.get("copia"),
                   r.get("archivo"),
                   "; ".join(r["incidencias"])])
        if r.get("estado") == "revisar":
            for celda in ws[ws.max_row]:
                celda.fill = amarillo
        ws.cell(ws.max_row, 3).number_format = "0.00"

    # ---------------------------------------------------------------- Lista de clase
    if alumnos:
        wl = wb.create_sheet(t("Lista de clase"))
        _cabecera(wl, ["Documento", "Apellidos y nombre", "Nota", "Observación"], [14, 36, 9, 40])
        por_doc = {}
        for r in res:
            if r.get("alumno"):
                por_doc.setdefault(r["alumno"]["documento"], []).append(r)
        for a in alumnos:
            hojas = por_doc.get(a.documento, [])
            if not hojas:
                wl.append([a.documento, a.nombre, t("NP"), t("No presentado")])
            else:
                obs = t("Varias hojas") if len(hojas) > 1 else (t("Pendiente de revisar")
                                                            if hojas[0]["estado"] == "revisar" else "")
                wl.append([a.documento, a.nombre, hojas[0].get("nota"), obs])
                wl.cell(wl.max_row, 3).number_format = "0.00"

    # ---------------------------------------------------------------- Respuestas
    n = gen["n_preguntas"]
    wr = wb.create_sheet(t("Respuestas"))
    _cabecera(wr, ["Documento", "Modelo"] + [t("P{n}", n=i) for i in range(1, n + 1)],
              [14, 8] + [5] * n)
    for r in res:
        if not r.get("preguntas"):
            continue
        wr.append([(r.get("alumno") or {}).get("documento", r.get("id_leido")), r.get("modelo_nombre")]
                  + [p["marcadas"] + ("*" if p["estado"] == "revisar" else "") for p in r["preguntas"]])
    wr.append([])
    wr.append(["* = casilla «Anular» marcada pendiente de revisar"])

    # ---------------------------------------------------------------- Puntos por pregunta del banco
    ids = list(gen["banco"].keys())
    usados = [i for i in ids if any(i == p["id"] for r in res for p in r.get("preguntas", []))
              or any(i in (r.get("superadas") or []) for r in res)]
    wp = wb.create_sheet(t("Puntos por pregunta"))
    _cabecera(wp, ["Documento", "Modelo"] + [t("Id {n}", n=i) for i in usados], [14, 8] + [6] * len(usados))
    for r in res:
        if not r.get("preguntas"):
            continue
        pts = {p["id"]: (None if p["estado"] == "anulada" else p["puntos"]) for p in r["preguntas"]}
        pts.update({q: float(cfg["valor_bien"]) for q in r.get("superadas") or []})
        wp.append([(r.get("alumno") or {}).get("documento", r.get("id_leido")), r.get("modelo_nombre")]
                  + [pts.get(i) for i in usados])

    # ---------------------------------------------------------------- Estadística
    we = wb.create_sheet(t("Estadística"))
    _cabecera(we, ["Id", "Pregunta", "Correcta", "Veces", "% acierto", "Media", "Bien", "Parcial", "Mal",
                   "Blanco", "Eligen A", "Eligen B", "Eligen C", "Eligen D", "Observación"],
              [7, 60, 9, 7, 9, 7, 6, 7, 6, 7, 8, 8, 8, 8, 30])
    for e in estadistica(proyecto, res):
        obs = []
        if e["anulada"]:
            obs.append(t("Anulada"))
        if e["clave_cambiada"]:
            obs.append(t("Clave corregida (antes {x})", x=e["correctas_banco"]))
        we.append([e["id"], e["texto"], e["correctas"], e["n"], e["acierto"], e["media"], e["bien"],
                   e["parcial"], e["mal"], e["blanco"], *e["elecciones"], "; ".join(obs)])

    # ---------------------------------------------------------------- Configuración
    wc = wb.create_sheet(t("Configuración"))
    _cabecera(wc, ["Parámetro", "Valor"], [32, 60])
    for k, v in [("Examen", cfg["titulo"]), ("Asignatura", cfg.get("asignatura", "")),
                 ("Convocatoria", t(cfg.get("convocatoria", "")) if cfg.get("convocatoria") else ""), ("Curso", cfg.get("curso", "")), ("Código", gen["codigo"]), ("Generado", gen["fecha"]),
                 ("Exportado", time.strftime("%Y-%m-%d %H:%M")), ("Preguntas", n),
                 ("Valor acierto", cfg["valor_bien"]), ("Valor fallo", cfg["valor_mal"]),
                 ("Multirrespuesta", t("Puntuación parcial") if cfg["multiples"] == "parcial" else t("Solo cuenta si es exacta")), ("Nota máxima", cfg["nota_maxima"]),
                 ("Nota negativa permitida", t("sí") if cfg["permitir_negativa"] else t("no")),
                 ("Preguntas anuladas", ", ".join(cfg.get("anuladas") or []) or t("ninguna")),
                 ("Claves corregidas", ", ".join(f"{k}→{v}" for k, v in (cfg.get("claves") or {}).items())
                  or t("ninguna"))]:
        wc.append([t(k), v])

    for hoja in wb.worksheets:
        _fuente(hoja)
    destino = destino or proyecto.ruta("notas.xlsx")
    wb.save(destino)
    return str(destino)
